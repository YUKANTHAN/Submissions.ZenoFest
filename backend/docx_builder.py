"""
Branded DOCX generation for ZenoFest event submissions.

Every submission is appended as a fresh page to a single shared DOCX file:
  1. Read the existing file (or start a new one).
  2. Append a branded section per submission.
  3. Return the new bytes so the caller can upload them to Drive.
"""
import io
from datetime import datetime

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches

ZENOFEST_PURPLE = RGBColor(0x6A, 0x0D, 0xAD)
ZENOFEST_CYAN = RGBColor(0x00, 0x91, 0xEA)
DARK_GRAY = RGBColor(0x33, 0x33, 0x33)


def _add_header_rule(doc, color=ZENOFEST_CYAN):
    """Thick colored horizontal rule under a heading."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("____________________________________________")
    run.font.color.rgb = color
    run.font.size = Pt(8)
    return p


def _add_field(doc, label, value):
    p = doc.add_paragraph()
    run = p.add_run(f"{label}: ")
    run.bold = True
    run.font.color.rgb = DARK_GRAY
    run.font.size = Pt(11)
    value_run = p.add_run(str(value) if value not in (None, "") else "—")
    value_run.font.color.rgb = DARK_GRAY
    value_run.font.size = Pt(11)
    return p


def build_initial_doc():
    """Create a fresh populated DOCX with a title page (no submission yet)."""
    doc = Document()

    # Title
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("ZENOFEST")
    run.bold = True
    run.font.size = Pt(40)
    run.font.name = "Georgia"
    run.font.color.rgb = ZENOFEST_PURPLE
    _add_header_rule(doc, ZENOFEST_PURPLE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("2026 — Event Data Submissions")
    run.font.size = Pt(18)
    run.font.color.rgb = ZENOFEST_CYAN
    run.font.name = "Georgia"
    run.italic = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Official record of tech-event submissions from registered teams.")
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_GRAY

    doc.add_paragraph()
    return doc


def append_submission(existing_bytes, submission):
    """Append one page for a submission.
    `existing_bytes` is the current DOCX content or None to start fresh.
    `submission` is a dict with keys:
        team_id, team_name, tech_event, leader_name, college, email, submitted_at,
        form_data (dict of field_key -> value)
        field_labels (dict of field_key -> display label)
    Returns new DOCX bytes."""
    if existing_bytes:
        doc = Document(io.BytesIO(existing_bytes))
    else:
        doc = build_initial_doc()

    doc.add_page_break()

    # Event header
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Event Data Submission")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = ZENOFEST_PURPLE
    run.font.name = "Georgia"
    _add_header_rule(doc)

    # Team block
    _add_field(doc, "Tech Event", submission.get("tech_event", ""))
    if submission.get("team_id"):
        _add_field(doc, "Team ID", submission["team_id"])
    _add_field(doc, "Team Name", submission.get("team_name", ""))
    if submission.get("leader_name"):
        _add_field(doc, "Leader Name", submission["leader_name"])
    if submission.get("college"):
        _add_field(doc, "College", submission["college"])
    if submission.get("email"):
        _add_field(doc, "Email", submission["email"])
    _add_field(doc, "Submitted At", submission.get("submitted_at", ""))

    doc.add_paragraph()

    # Section divider
    p = doc.add_paragraph()
    run = p.add_run("——— Submission Data ———")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = ZENOFEST_CYAN

    # Form fields
    labels = submission.get("field_labels") or {}
    for key, value in (submission.get("form_data") or {}).items():
        label = labels.get(key, key.replace("_", " ").title())
        _add_field(doc, label, value)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()